# -*- coding: utf-8 -*-
"""
taogiong_chia_cau.py — Xuất file XEM TRƯỚC cách chia câu của kịch bản trước khi tạo giọng.

Đi đúng đường của lúc tạo giọng thật (_batch_run_tts trong amain_taogiong_gui.py và
web/runners/run_tts.py): đọc input.docx (hay input.txt tập cũ) → clean_text → chữ thường
→ split_chunks(chunk). Mỗi dòng trong file là MỘT đoạn sẽ đưa vào OmniVoice, kèm:
    số thứ tự | số ký tự | thời lượng dự kiến (vn_duration) | nghỉ chèn SAU đoạn | nội dung
Nghỉ sau đoạn: ¶ 0,80 s (hết đoạn văn) · . 0,45 s (hết câu) · , 0,22 s (hết vế ; : ,).

Kết quả: input_chia_cau.docx nằm cạnh file input. Lúc tạo giọng thật cũng ghi lại file này
(gọi write_preview) nên nó luôn khớp với bản audio vừa sinh.

Chạy:
    python taogiong_chia_cau.py --episode 102            # tập trong kịch_bản/
    python taogiong_chia_cau.py --episode 102 --chunk 200
    python taogiong_chia_cau.py --input "duong_dan/input.docx"
"""

import sys
import os

_REPO_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), os.pardir, os.pardir))
_VENV_PYTHON = os.path.join(_REPO_ROOT, "venv", "Scripts", "python.exe")
if __name__ == "__main__" and os.path.exists(_VENV_PYTHON) and \
        os.path.normcase(os.path.abspath(sys.executable)) != \
        os.path.normcase(os.path.abspath(_VENV_PYTHON)):
    import subprocess
    sys.exit(subprocess.run([_VENV_PYTHON] + sys.argv).returncode)

try:
    sys.stdout.reconfigure(encoding="utf-8", errors="replace")
except Exception:
    pass

_SCRIPTS_DIR = os.path.dirname(os.path.abspath(__file__))
_BASE_DIR = os.path.dirname(_SCRIPTS_DIR)
for _p in (_REPO_ROOT, _SCRIPTS_DIR, os.path.join(_BASE_DIR, "YOUTUBE"), _BASE_DIR):
    if _p not in sys.path:
        sys.path.insert(0, _p)

import argparse
import json
from pathlib import Path

PREVIEW_NAME = "input_chia_cau.docx"
OPTS_FILE = Path(_BASE_DIR) / "taogiong_options.json"


def _gui():
    import amain_taogiong_gui as gui
    return gui


def chunk_size_from_options(default=160) -> int:
    """Ô 'chunk' đang lưu trong taogiong_options.json (GUI và web dùng chung)."""
    try:
        return int(json.loads(OPTS_FILE.read_text(encoding="utf-8")).get("chunk", default))
    except Exception:
        return default


def split_for_tts(text: str, chunk: int) -> list:
    """Đúng ba bước của lúc tạo giọng: clean_text → lower → split_chunks."""
    gui = _gui()
    return gui.split_chunks(gui.clean_text(text).lower(), chunk)


def _pause_label(gui, chunk_text: str) -> tuple[str, float]:
    sec = gui._pause_after(chunk_text)
    if sec == gui.PAUSE_PARA:
        return "¶", sec
    if sec == gui.PAUSE_SENT:
        return ".", sec
    return ",", sec


def describe(chunks: list) -> dict:
    """Số liệu tổng: số đoạn, thời lượng dự kiến, phân bố nghỉ, đoạn dài/ngắn nhất."""
    gui = _gui()
    rows, total_speech, total_pause = [], 0.0, 0.0
    n_para = n_sent = n_clause = 0
    for i, c in enumerate(chunks, 1):
        text = " ".join(c.split())
        dur = gui.vn_duration(text) or 0.0
        mark, psec = _pause_label(gui, c)
        if i < len(chunks):
            total_pause += psec
        if mark == "¶":
            n_para += 1
        elif mark == ".":
            n_sent += 1
        else:
            n_clause += 1
        total_speech += dur
        rows.append((i, len(text), dur, mark, psec, text))
    lens = [r[1] for r in rows] or [0]
    return {
        "rows": rows, "n": len(rows),
        "speech_sec": total_speech, "pause_sec": total_pause,
        "n_para": n_para, "n_sent": n_sent, "n_clause": n_clause,
        "min_len": min(lens), "max_len": max(lens),
        "avg_len": (sum(lens) / len(lens)) if rows else 0,
    }


def write_preview(folder_or_input, chunks: list, chunk_size: int, source_name: str = "") -> Path:
    """Ghi input_chia_cau.docx cạnh file input. folder_or_input: thư mục tập hoặc đường
    dẫn file input (lấy thư mục cha)."""
    from docx import Document
    from docx.shared import Pt, RGBColor

    p = Path(folder_or_input)
    folder = p if p.is_dir() else p.parent
    out = folder / PREVIEW_NAME
    d = describe(chunks)
    gui = _gui()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Consolas"
    style.font.size = Pt(9.5)

    doc.add_heading(f"Chia câu để tạo giọng — {source_name or folder.name}", level=1)
    tong = d["speech_sec"] + d["pause_sec"]
    doc.add_paragraph(
        f"{d['n']} đoạn (mỗi đoạn một câu; ngưỡng cắt câu dài = {chunk_size} ký tự) · "
        f"ngắn nhất {d['min_len']} · dài nhất {d['max_len']} · trung bình {d['avg_len']:.0f} ký tự")
    doc.add_paragraph(
        f"Thời lượng dự kiến ≈ {tong / 60:.1f} phút "
        f"(đọc {d['speech_sec'] / 60:.1f} phút + nghỉ {d['pause_sec'] / 60:.1f} phút). "
        f"Nghỉ sau đoạn: ¶ {gui.PAUSE_PARA:.2f} s hết đoạn văn ×{d['n_para']} · "
        f". {gui.PAUSE_SENT:.2f} s hết câu ×{d['n_sent']} · "
        f", {gui.PAUSE_CLAUSE:.2f} s hết vế ×{d['n_clause']}")
    doc.add_paragraph("Cột: số thứ tự | số ký tự | giây dự kiến | nghỉ sau | nội dung đưa vào model "
                      "(đã hạ chữ thường, bỏ dấu lạ như lúc tạo giọng thật).")

    grey = RGBColor(0x88, 0x88, 0x88)
    for i, n_chars, dur, mark, psec, text in d["rows"]:
        para = doc.add_paragraph()
        head = para.add_run(f"{i:03d} | {n_chars:3d} | {dur:5.1f}s | {mark} {psec:.2f} | ")
        head.font.color.rgb = grey
        para.add_run(text)
        para.paragraph_format.space_after = Pt(1)
    doc.save(str(out))
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Xuất file xem trước cách chia câu (input_chia_cau.docx).")
    ap.add_argument("--episode", default="", help="Số tập trong kịch_bản/.")
    ap.add_argument("--input", default="", help="Đường dẫn file input.docx/.txt (thay cho --episode).")
    ap.add_argument("--chunk", type=int, default=0,
                    help="Ngưỡng cắt câu dài (mặc định: ô 'chunk' trong taogiong_options.json).")
    args = ap.parse_args(argv)

    import dich_input_docx as inputdocx
    gui = _gui()
    if args.input:
        src = inputdocx.resolve_input(args.input)
    elif args.episode:
        folder = gui.find_episode_dir(args.episode)
        if folder is None:
            print(f"❌ Không tìm thấy tập {args.episode} trong {gui.SCRIPT_DIR}")
            return 2
        src = inputdocx.find_input(folder)
        if src is None:
            print(f"❌ Tập {args.episode} chưa có input.docx / input.txt.")
            return 2
    else:
        ap.error("cần --episode hoặc --input")
    if not src.is_file():
        print(f"❌ Không thấy file: {src}")
        return 2

    chunk = args.chunk or chunk_size_from_options()
    text = inputdocx.read_input_text(src)
    chunks = split_for_tts(text, chunk)
    if not chunks:
        print(f"❌ {src.name} không có nội dung.")
        return 2
    out = write_preview(src, chunks, chunk, src.name)
    d = describe(chunks)
    tong = d["speech_sec"] + d["pause_sec"]
    print(f"✅ {d['n']} đoạn · dự kiến ≈ {tong / 60:.1f} phút · ngắn nhất {d['min_len']} · "
          f"dài nhất {d['max_len']} ký tự (ngưỡng {chunk}) → {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
